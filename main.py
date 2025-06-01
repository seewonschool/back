from fastapi import FastAPI
from fastapi.responses import FileResponse
from typing import Union
from pydantic import BaseModel

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl import load_workbook
from const.kakao_conversation import KakaoConversaionId, MajorCode
from model.major import Major
import requests
import firebase_admin
from firebase_admin import credentials
from firebase_admin import firestore
from datetime import datetime
import os

app = FastAPI()

# firebase 계정 키 파일 경로
cred = credentials.Certificate("firebasekey.json")
firebase_admin.initialize_app(cred)
# Firestore 클라이언트 초기화
db = firestore.client()

headers = {
        'Authorization': 'Bearer e0530d2e.12f54b7fdb1f4628a6280c34eee7ef0c',
        'Content-Type': 'application/json'
}

### 문서 데이터 추출 -> 엑셀을 만들어서 -> 삭제 -> 등록 -> 초대

@app.get("/users/firebase")
def get_students_collection():
    # 'students' 컬렉션의 모든 문서 가져오기
    students_ref = db.collection('students')
    docs = students_ref.stream()

    # 2. 엑셀 파일 생성
    wb = Workbook()  # 새 엑셀 파일 생성
    wb._write_version = "2210"                 # fileVersion.appName 대신 엑셀 버전
    ws = wb.active   # 기본 워크시트 선택
    # ws.title = "Students Data"  # 워크시트 이름 설정
    # 메타데이터 설정
    wb.properties.title = "Example Excel File"
    wb.properties.creator = "My Application"
    wb.properties.description = "This is a test Excel file created with Python."
    wb.properties.lastModifiedBy = "AutomationBot"
    wb.properties.created        = datetime.utcnow()
    wb.properties.modified       = datetime.utcnow()
    ws.title = "Sheet1"                        # ① 샘플과 동일한 시트명
    ws.sheet_view.tabSelected = True
    

    # 데이터 헤더 추가
    headers = [ "ID(이메일)*",
                "이름*",
                "닉네임",
                "조직",
                "직책",
                "직위",
                "휴대전화",
                "유선전화",
                "이메일",
                "근무위치",
                "생일",
                "상태 메세지"]  # 각 필드에 맞는 헤더 설정
    ws.append(headers)  # 첫 번째 행에 헤더 추가

    for doc in docs :
      row = {"id": doc.id, **doc.to_dict()}
      # 엑셀에 데이터 추가
      ws.append([row.get("email", ""), row.get("name", ""), "", row.get("major", ""),"", "","","",row.get("email", ""),"", "", ""])

    
    # 엑셀파일 폰트 설정 ...
    font_style = Font(name="Malgun Gothic", size=11)
    for row in ws.iter_rows():  # 시트의 모든 행을 순회
        for cell in row:  # 해당 행의 모든 셀을 순회
            cell.font = font_style

    # 3. 파일 저장 경로 설정
    file_name = "students_data"
    wb.save(file_name)  # 엑셀 파일 저장

    # 4. firebase 내 데이터 삭제
    students_ref_re = db.collection('students') #-> 중간에 들어온 사람은 엑셀 못들어가고 삭제될듯..ㅎ?
    docs_re = students_ref_re.stream()
    for doc in docs_re:
        doc.reference.delete()
    
    # 5. 파일 응답 반환
    # 파일이 존재하는지 확인
    if not os.path.exists(file_name):
        return {"error": "File not found!"}
    
    # 올바른 파일 경로로 반환
    return FileResponse(
        path=file_name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=file_name
    )

@app.get("/users/kakao")
def get_kakao_users():
    res = requests.get("https://api.kakaowork.com/v1/users.list", headers=headers)
    result = res.json()
    return result["users"]

@app.get("/room/invite")
def invite_room():
    users = get_kakao_users()
    response = []
    for user in users:
      print(user["name"], user["department"])
      if(MajorCode.get(user["department"]) == None): continue
      print(user["name"], MajorCode.get(user["department"]),KakaoConversaionId[MajorCode.get(user["department"])].value)
      data = {
          'user_ids': [user["id"]]
      }
      res = requests.post(f"https://api.kakaowork.com/v1/conversations/{KakaoConversaionId[MajorCode.get(user["department"])].value}/invite", headers=headers, json=data)
      result = res.json()
      response.append(result)
    return response

class KickUser(BaseModel):
    conversation_id: int
    user_ids: list[int]
@app.delete("/room/kick")
def kick_room(body: KickUser):
    data = {
        'user_ids': body.user_ids
    }
    res = requests.post(f"https://api.kakaowork.com/v1/conversations/{body.conversation_id}/kick", headers=headers, json=data)
    result = res.json()
    return result

@app.get("/conversation")
def make_chat():
    data = {
        'user_ids': [],
        'conversation_name': "빈채팅방"
    }
    res = requests.post("https://api.kakaowork.com/v1/conversations.open", headers=headers, json=data)
    result = res.json()
    return result


class Msg(BaseModel):
    conversation_id: int
    title: str
    date: str
    link: str

@app.post("/chat", summary="채팅 전송")
def send_chat(msg: Msg):
    data = {
            "conversation_id": msg.conversation_id,
            "text": msg.title,
            "blocks":[
                {
                  "type": "header",
                  "text": "공지사항",
                  "style": "yellow"
                },
                {
                  "type": "text",
                  "text": "text sample",
                  "inlines": [
                    {
                      "type": "styled",
                      "text": msg.title,
                      "bold": True
                    }
                  ]
                },
                {
                  "type": "description",
                  "term": "등록일",
                  "content": {
                    "type": "text",
                    "text": msg.date
                  },
                  "accent": True
                },
                {
                  "type": "button",
                  "text": "바로가기",
                  "style": "default",
                  "action": {
                    "type": "open_system_browser",
                    "name": "button1",
                    "value": msg.link
                  }
                }
              ]}
    res = requests.post("https://api.kakaowork.com/v1/messages.send", headers=headers, json=data)
    result = res.json()
    return result